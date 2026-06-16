from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = ROOT.parent
RESULT_DIR = ROOT / "experiments" / "results"
CHATGPT_XLSX = PROJECT_ROOT / "thực nghiệm 2(chatgpt).xlsx"
RUNTIME_CSV = RESULT_DIR / "experiment_images_local_gemini.csv"
OUTPUT_XLSX = PROJECT_ROOT / "thuc_nghiem_tong_hop_chatgpt_gemini_system_v2.xlsx"

TRANS = str.maketrans({
    "绪": "緒",
    "统": "統",
    "历": "曆",
    "万": "萬",
    "制": "製",
    "内": "內",
})


def norm_cjk(value) -> str:
    return "".join(ch for ch in str(value or "") if "\u4e00" <= ch <= "\u9fff").translate(TRANS)


def agree(a, b) -> bool:
    aa = norm_cjk(a)
    bb = norm_cjk(b)
    return bool(aa and bb and aa == bb)


def status(value) -> str:
    return "available" if norm_cjk(value) else "missing"


def pct(num: int, den: int) -> str:
    return f"{num / den:.2%}" if den else "n/a"


def main() -> None:
    chatgpt = pd.read_excel(CHATGPT_XLSX)
    runtime = pd.read_csv(RUNTIME_CSV, encoding="utf-8-sig")

    chatgpt = chatgpt.rename(columns={
        "file_name": "image",
        "visible_text_raw": "chatgpt_visible_text_raw",
        "chu_han_best": "chatgpt_text",
        "candidates": "chatgpt_candidates",
        "confidence": "chatgpt_confidence",
        "notes": "chatgpt_notes",
    })

    keep_runtime = [
        "image",
        "gemini_text",
        "gemini_confidence",
        "gemini_model",
        "gemini_status",
        "gemini_error",
        "system_text",
        "system_confidence",
        "system_aggregated_confidence",
        "system_primary_pipeline",
        "system_valid_pipelines",
        "system_latency",
        "system_status",
        "system_error",
        "system_gemini_model",
    ]
    for col in keep_runtime:
        if col not in runtime.columns:
            runtime[col] = ""
    runtime = runtime[keep_runtime]

    merged = chatgpt.merge(runtime, on="image", how="left")
    for col in ["chatgpt_text", "gemini_text", "system_text"]:
        merged[f"{col}_norm"] = merged[col].map(norm_cjk)
    merged["chatgpt_gemini_match"] = merged.apply(lambda r: agree(r["chatgpt_text"], r["gemini_text"]), axis=1)
    merged["chatgpt_system_match"] = merged.apply(lambda r: agree(r["chatgpt_text"], r["system_text"]), axis=1)
    merged["gemini_system_match"] = merged.apply(lambda r: agree(r["gemini_text"], r["system_text"]), axis=1)
    merged["chatgpt_status"] = merged["chatgpt_text"].map(status)
    merged["gemini_available"] = merged["gemini_text"].map(status)
    merged["system_available"] = merged["system_text"].map(status)

    total = len(merged)
    gemini_avail = int((merged["gemini_available"] == "available").sum())
    system_avail = int((merged["system_available"] == "available").sum())
    cg_den = int(((merged["chatgpt_status"] == "available") & (merged["gemini_available"] == "available")).sum())
    cs_den = int(((merged["chatgpt_status"] == "available") & (merged["system_available"] == "available")).sum())
    gs_den = int(((merged["gemini_available"] == "available") & (merged["system_available"] == "available")).sum())
    summary_rows = [
        ["Metric", "Value", "Note"],
        ["Total images in ChatGPT file", total, "Input workbook"],
        ["Gemini standalone available", gemini_avail, "From experiment_images_local_gemini.csv"],
        ["Full system voting available", system_avail, "Full system: OCR + OCR search + Gemini + OpenCode + image search + ML match + voting"],
        ["ChatGPT vs Gemini exact agreement", pct(int(merged["chatgpt_gemini_match"].sum()), cg_den), f"{int(merged['chatgpt_gemini_match'].sum())}/{cg_den} comparable rows"],
        ["ChatGPT vs System exact agreement", pct(int(merged["chatgpt_system_match"].sum()), cs_den), f"{int(merged['chatgpt_system_match'].sum())}/{cs_den} comparable rows"],
        ["Gemini vs System exact agreement", pct(int(merged["gemini_system_match"].sum()), gs_den), f"{int(merged['gemini_system_match'].sum())}/{gs_den} comparable rows"],
    ]

    method_rows = []
    if "system_primary_pipeline" in merged:
        method_rows.extend([["System primary pipeline", k, v] for k, v in merged["system_primary_pipeline"].fillna("").value_counts().items() if k])
    if "system_gemini_model" in merged:
        method_rows.extend([["System Gemini model", k, v] for k, v in merged["system_gemini_model"].fillna("").value_counts().items() if k])
    if "gemini_model" in merged:
        method_rows.extend([["Standalone Gemini model", k, v] for k, v in merged["gemini_model"].fillna("").value_counts().items() if k])
    method_df = pd.DataFrame(method_rows, columns=["Category", "Name", "Count"])

    output_cols = [
        "image",
        "chatgpt_text",
        "chatgpt_confidence",
        "chatgpt_notes",
        "gemini_text",
        "gemini_confidence",
        "gemini_model",
        "gemini_status",
        "system_text",
        "system_confidence",
        "system_aggregated_confidence",
        "system_primary_pipeline",
        "system_gemini_model",
        "system_latency",
        "system_status",
        "chatgpt_gemini_match",
        "chatgpt_system_match",
        "gemini_system_match",
        "system_error",
    ]
    results = merged[output_cols]

    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        pd.DataFrame(summary_rows[1:], columns=summary_rows[0]).to_excel(writer, sheet_name="Summary", index=False)
        results.to_excel(writer, sheet_name="Combined_Results", index=False)
        method_df.to_excel(writer, sheet_name="Method_Stats", index=False)
        chatgpt.to_excel(writer, sheet_name="ChatGPT_Raw", index=False)
        runtime.to_excel(writer, sheet_name="Runtime_Raw", index=False)

    wb = load_workbook(OUTPUT_XLSX)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
        for col_cells in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col_cells[:200])
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(max_len + 2, 10), 45)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws = wb["Summary"]
    ws["E1"] = "Agreement Chart"
    chart_data_start = 10
    chart_rows = [
        ["Pair", "Comparable", "Exact matches"],
        ["ChatGPT-Gemini", cg_den, int(merged["chatgpt_gemini_match"].sum())],
        ["ChatGPT-System", cs_den, int(merged["chatgpt_system_match"].sum())],
        ["Gemini-System", gs_den, int(merged["gemini_system_match"].sum())],
    ]
    for r_idx, row in enumerate(chart_rows, chart_data_start):
        for c_idx, value in enumerate(row, 5):
            ws.cell(r_idx, c_idx, value)
    chart = BarChart()
    chart.title = "Exact Agreement Counts"
    chart.y_axis.title = "Images"
    chart.x_axis.title = "Method Pair"
    chart.add_data(Reference(ws, min_col=7, min_row=chart_data_start, max_row=chart_data_start + 3), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=5, min_row=chart_data_start + 1, max_row=chart_data_start + 3))
    chart.height = 7
    chart.width = 12
    ws.add_chart(chart, "E3")

    wb.save(OUTPUT_XLSX)
    print(OUTPUT_XLSX)


if __name__ == "__main__":
    main()
